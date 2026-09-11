"""Native package behavior, using genuine editable PowerPoint fixtures."""
import argparse
from hashlib import sha256
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))
from inspect_pptx_package import inspect_package
from ppt_pipeline import (build_translation_manifest, command_apply, command_inspect,
                          command_prepare, command_verify, PipelineError)
from pptx_ooxml import apply_manifest, OoxmlError
from validate_manifest import ManifestError


class NativeOverlayTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source, self.output = self.root / "source.pptx", self.root / "output.pptx"
        self.manifest_path = self.root / "manifest.json"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(.1), Inches(4), Inches(.5)).text = "设备"
        image = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(image, "PNG")
        image.seek(0)
        self.picture = slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(4), Inches(3))
        presentation.save(self.source)
        self.manifest = build_translation_manifest(inspect_package(self.source), "en")
        self.manifest["translation_units"][0]["translation"] = "Equipment"
        self.overlay = {
            "id": "label-1", "kind": "office_overlay", "localization_mode": "bilingual_below",
            "source_text": "冷却机", "translation": "Cooler & fan\nMotor",
            "location": {"page_or_slide": 1, "host_shape_id": self.picture.shape_id, "region_id": "label-1"},
            "source_region": {"x": .1, "y": .1, "w": .5, "h": .1},
            "region": {"x": .1, "y": .21, "w": .5, "h": .2},
            "background": {"mode": "transparent"},
            "style": {"font_name": "Arial", "font_size_pt": 12, "bold": True,
                      "text_rgb": "123456", "fill_rgb": "FFFFFF", "align": "center"},
        }
        self.manifest["overlays"] = [self.overlay]
        self.manifest["image_groups"][0].update(decision="overlay", overlay_ids=["label-1"], preserve_source_image=True)

    def save_manifest(self):
        self.manifest_path.write_text(json.dumps(self.manifest), encoding="utf-8")

    def test_real_deck_overlays_editable_with_geometry_style_and_unchanged_host(self):
        self.save_manifest()
        apply_manifest(self.source, self.manifest_path, self.output)
        slide = Presentation(self.output).slides[0]
        boxes = [s for s in slide.shapes if s.name == "office-translate-overlay:label-1"]
        self.assertEqual(1, len(boxes), "native writer must insert an editable overlay")
        box = boxes[0]
        self.assertEqual(self.overlay["translation"], box.text)
        self.assertEqual(round(self.picture.left + .1 * self.picture.width), box.left)
        self.assertEqual(round(self.picture.top + .21 * self.picture.height), box.top)
        self.assertEqual(round(.5 * self.picture.width), box.width)
        self.assertEqual(round(.2 * self.picture.height), box.height)
        run = box.text_frame.paragraphs[0].runs[0]
        self.assertEqual("Arial", run.font.name)
        self.assertEqual(12, run.font.size.pt)
        self.assertTrue(run.font.bold)
        self.assertEqual("123456", str(run.font.color.rgb))
        self.assertTrue(box._element.xpath("./p:spPr/a:noFill"))
        self.assertEqual("Equipment", slide.shapes[0].text)
        with ZipFile(self.source) as before, ZipFile(self.output) as after:
            for name in before.namelist():
                if name != "ppt/slides/slide1.xml":
                    self.assertEqual(before.read(name), after.read(name), name)
            import lxml.etree as ET
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            original = ET.fromstring(before.read("ppt/slides/slide1.xml")).find(".//p:pic", ns)
            result = ET.fromstring(after.read("ppt/slides/slide1.xml")).find(".//p:pic", ns)
            self.assertEqual(ET.tostring(original), ET.tostring(result))

    def test_failed_late_text_write_keeps_prior_output(self):
        self.manifest["occurrences"][0]["source_text"] = "wrong"
        self.manifest["translation_units"][0]["source_text"] = "wrong"
        self.save_manifest()
        self.output.write_bytes(b"previous accepted output")
        with self.assertRaises(OoxmlError):
            apply_manifest(self.source, self.manifest_path, self.output)
        self.assertEqual(b"previous accepted output", self.output.read_bytes())

    def test_invalid_xml_translation_keeps_prior_output(self):
        self.manifest["overlays"] = []
        self.manifest["translation_units"][0]["translation"] = "Equipment\x00"
        self.save_manifest()
        self.output.write_bytes(b"previous accepted output")
        with self.assertRaises((OoxmlError, ValueError)):
            apply_manifest(self.source, self.manifest_path, self.output)
        self.assertEqual(b"previous accepted output", self.output.read_bytes())

    def test_rotated_host_fails_with_overlay_id_and_keeps_prior_output(self):
        presentation = Presentation(self.source)
        presentation.slides[0].shapes[1].rotation = 30
        presentation.save(self.source)
        self.manifest["source_sha256"] = sha256(self.source.read_bytes()).hexdigest()
        self.save_manifest()
        self.output.write_bytes(b"previous accepted output")
        with self.assertRaisesRegex(OoxmlError, "label-1.*rotat"):
            apply_manifest(self.source, self.manifest_path, self.output)
        self.assertEqual(b"previous accepted output", self.output.read_bytes())

    def test_complex_parts_are_disclosed_preserved_and_verified(self):
        part = "ppt/charts/chart99.xml"
        with ZipFile(self.source, "a") as package:
            package.writestr(part, '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:v>Sales</c:v></c:chartSpace>')
        inventory = inspect_package(self.source)
        self.assertEqual(part, inventory["preserved_parts"][0]["part"])
        self.assertIn("untranslated", inventory["warnings"][0]["message"])
        self.manifest = build_translation_manifest(inventory, "en")
        self.manifest["translation_units"][0]["translation"] = "Equipment"
        self.assertEqual(inventory["preserved_parts"], self.manifest["preserved_parts"])
        self.save_manifest()
        apply_manifest(self.source, self.manifest_path, self.output)
        with ZipFile(self.source) as before, ZipFile(self.output) as after:
            self.assertEqual(before.read(part), after.read(part))

        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        (job / "translation-manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        with ZipFile(self.output) as package:
            entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
        with ZipFile(self.output, "w") as package:
            for entry, payload in entries:
                package.writestr(entry, payload.replace(b"Sales", b"Changed") if entry.filename == part else payload)
        self.assertEqual(2, command_verify(argparse.Namespace(source=self.source, job_dir=job, output=self.output)))
        errors = json.loads((job / "verification.json").read_text(encoding="utf-8"))["errors"]
        self.assertIn({"code": "preserved-part-changed", "part": part}, errors)

    def test_zip_write_error_keeps_prior_output(self):
        self.save_manifest()
        self.output.write_bytes(b"previous accepted output")
        with patch("pptx_ooxml.zipfile.ZipFile.writestr", side_effect=OSError("disk full")):
            with self.assertRaisesRegex(OSError, "disk full"):
                apply_manifest(self.source, self.manifest_path, self.output)
        self.assertEqual(b"previous accepted output", self.output.read_bytes())
        self.assertEqual([], list(self.root.glob(".output.pptx.*.tmp")))

    def test_grouped_host_fails_with_actionable_overlay_error(self):
        presentation = Presentation(self.source)
        slide = presentation.slides[0]
        slide.shapes.add_group_shape([slide.shapes[1]])
        presentation.save(self.source)
        self.save_manifest()
        with self.assertRaisesRegex(OoxmlError, "label-1.*grouped"):
            apply_manifest(self.source, self.manifest_path, self.output)

    def test_converted_working_copy_cannot_be_overwritten_or_changed(self):
        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        (job / "translation-manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        working = self.root / "working-source.pptx"
        working.write_bytes(self.source.read_bytes())
        inventory_path = job / "inventory.json"
        inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
        inventory["working_source_path"] = str(working)
        inventory_path.write_text(json.dumps(inventory), encoding="utf-8")
        with self.assertRaisesRegex(PipelineError, "overwrite the working source"):
            command_apply(argparse.Namespace(input=self.source, job_dir=job, output=working))
        with ZipFile(working, "a") as package:
            package.writestr("changed.txt", "mutation")
        with self.assertRaisesRegex(PipelineError, "working source hash changed"):
            command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output))

    def test_pipeline_overlay_apply_never_starts_com(self):
        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        command_prepare(argparse.Namespace(job_dir=job, source_language="auto"))
        (job / "translation-manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        with patch("ppt_pipeline.subprocess.run", side_effect=AssertionError("COM apply is forbidden")):
            command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output))
        self.assertEqual(0, command_verify(argparse.Namespace(source=self.source, job_dir=job, output=self.output)))

    def test_prepare_apply_verify_accepts_equivalent_technical_unit_spacing(self):
        presentation = Presentation(self.source)
        presentation.slides[0].shapes[0].text = "电机 45kW 温度 25℃"
        presentation.save(self.source)
        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        command_prepare(argparse.Namespace(job_dir=job, source_language="auto"))
        path = job / "translation-manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest["translation_units"][0]["translation"] = "Motor 45 kW Temperature 25 °C"
        manifest["image_groups"][0]["decision"] = "skip_unclear"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        self.assertEqual(0, command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output)))
        self.assertEqual(0, command_verify(argparse.Namespace(source=self.source, job_dir=job, output=self.output)))
        accepted = self.output.read_bytes()
        manifest["translation_units"][0]["translation"] = "Motor 46 kW Temperature 25 °C"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaises(ManifestError):
            command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output))
        self.assertEqual(accepted, self.output.read_bytes())

    def test_explicit_custom_token_is_checked_during_apply_and_verify(self):
        presentation = Presentation(self.source)
        presentation.slides[0].shapes[0].text = "AGICO 设备"
        presentation.save(self.source)
        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        command_prepare(argparse.Namespace(job_dir=job, source_language="auto"))
        path = job / "translation-manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest["translation_units"][0]["protected_tokens"] = ["AGICO"]
        manifest["occurrences"][0]["protected_tokens"] = ["AGICO"]
        manifest["image_groups"][0]["decision"] = "skip_unclear"
        for translation in ("Equipment", "Agico Equipment", "XAGICOY Equipment", "A G I C O Equipment"):
            with self.subTest(translation=translation):
                manifest["translation_units"][0]["translation"] = translation
                path.write_text(json.dumps(manifest), encoding="utf-8")
                self.output.write_bytes(b"previous accepted output")
                with self.assertRaisesRegex(ManifestError, "protected token"):
                    command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output))
                self.assertEqual(b"previous accepted output", self.output.read_bytes())
                # Simulate a file written externally; verify must catch the same damage
                # even when the erroneous manifest translation matches output text.
                apply_manifest(self.source, path, self.output)
                self.assertEqual(2, command_verify(argparse.Namespace(source=self.source, job_dir=job, output=self.output)))
                errors = json.loads((job / "verification.json").read_text(encoding="utf-8"))["errors"]
                self.assertIn({"code": "technical-parameter-mismatch", "id": manifest["occurrences"][0]["id"]}, errors)
        manifest["translation_units"][0]["translation"] = "AGICO Equipment"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        self.assertEqual(0, command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output)))
        self.assertEqual(0, command_verify(argparse.Namespace(source=self.source, job_dir=job, output=self.output)))

    def test_pipeline_rejects_changed_source_before_output_mutation(self):
        job = self.root / "job"
        command_inspect(argparse.Namespace(input=self.source, job_dir=job, target_language="en"))
        (job / "translation-manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        with ZipFile(self.source, "a") as package:
            package.writestr("changed.txt", "mutation")
        self.output.write_bytes(b"previous accepted output")
        with self.assertRaisesRegex((PipelineError, OoxmlError), "source.*(changed|hash)"):
            command_apply(argparse.Namespace(input=self.source, job_dir=job, output=self.output))
        self.assertEqual(b"previous accepted output", self.output.read_bytes())


if __name__ == "__main__":
    unittest.main()
