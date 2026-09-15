"""Paired rows built from the original package, retaining source styles and formulas."""
from copy import deepcopy
from functools import lru_cache
import math
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter
from lxml import etree as ET
from PIL import ImageFont

from excel_native_ooxml import (Package, Q, XML_SPACE, canonical, coordinates,
                               decisions, layout_geometry, serialize, xml)


@lru_cache(maxsize=32)
def translation_font(size, bold):
    name = 'arialbi.ttf' if bold else 'ariali.ttf'
    font_path = Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/name
    try:
        return ImageFont.truetype(str(font_path),round(size*96/72))
    except OSError:
        return None


def translated_height(text, width, size, bold=False):
    font = translation_font(size,bold)
    pixels = max(7,width*7-6)  # Cell padding; use conservative Excel digit width.
    measure = font.getlength if font else lambda s: sum(size*(0.8 if ord(c)<128 else 1.4) for c in s)
    lines = 0
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = (line+' '+word).strip()
            if line and measure(candidate) > pixels:
                lines += 1
                line = ''
            for character in ((' ' if line else '')+word):
                if line and measure(line+character)>pixels:
                    lines += 1
                    line = ''
                line += character
        lines += 1
    return math.ceil(lines*size*1.5+6)


def address(value, translation=False):
    col, row = coordinates(value)
    if row * 2 > 1048576:
        raise ValueError('bilingual row count exceeds Excel limit')
    return f'{get_column_letter(col)}{row * 2 - (0 if translation else 1)}'


def map_reference(value, include_translation=False):
    # Quoted sheet names are opaque. Unquoted names must not be mistaken for cells.
    prefix, sep, ref = value.rpartition('!')
    if not sep:
        ref = value
    prefix = prefix + sep
    parts = ref.split(':')
    out = []
    for index, part in enumerate(parts):
        match = re.fullmatch(r'(\$?[A-Za-z]{1,3})?(\$?)(\d+)', part)
        if match:
            row = int(match[3]) * 2 - 1
            if include_translation and (index == 1 or len(parts) == 1):
                row += 1
            if row > 1048576:
                raise ValueError('bilingual reference exceeds Excel limit')
            out.append(f'{match[1] or ""}{match[2]}{row}')
        else:
            out.append(part)
    return prefix + ':'.join(out)


def map_formula(formula):
    tokens = Tokenizer('=' + formula.lstrip('=')).items
    sensitive = {'ROW','ROWS','COLUMN','COLUMNS','OFFSET','INDIRECT','ADDRESS','CELL',
                 'COUNTA','COUNTBLANK','COUNTIF','COUNTIFS','SUMIF','SUMIFS','AVERAGEIF',
                 'AVERAGEIFS','SUMPRODUCT','INDEX','MATCH','XMATCH','XLOOKUP','VLOOKUP',
                 'HLOOKUP','FILTER','SORT','UNIQUE','SEQUENCE','TEXTJOIN','CONCAT','CONCATENATE'}
    for token in tokens:
        if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value[:-1].upper() in sensitive:
            raise ValueError(f'unsupported bilingual formula function: {token.value}')
        if token.type == 'OPERAND' and token.subtype == 'RANGE':
            if '[' in token.value or re.search(r'[^\s!]:[^\s!]+!', token.value):
                raise ValueError(f'unsupported bilingual reference: {token.value}')
            token.value = map_reference(token.value)
    return ''.join(t.value for t in tokens)


def check_source(package):
    for sheet in package.sheets:
        root = sheet['root']
        for node in root.findall(f"{Q('mergeCells')}/{Q('mergeCell')}"):
            first, last = node.get('ref').split(':')
            if coordinates(first)[1] != coordinates(last)[1]:
                raise ValueError(f'unsupported bilingual vertical merge: {node.get("ref")}')
        # These features need a dedicated row-aware writer; do not silently lose them.
        for name in ('drawing','legacyDrawing','tableParts','autoFilter','extLst','conditionalFormatting'):
            if root.find(Q(name)) is not None:
                raise ValueError(f'unsupported bilingual row-sensitive object: {name}')
    for _, cell in package.cells.items():
        node = cell.find(Q('f'))
        if node is not None:
            if node.get('t') not in {None, 'normal', 'shared'}:
                raise ValueError(f'unsupported bilingual formula type: {node.get("t")}')
            if node.text:
                map_formula(node.text)


class TranslationStyles:
    def __init__(self, styles):
        self.styles = styles
        self.xfs = styles.find(Q('cellXfs'))
        self.fonts = styles.find(Q('fonts'))
        fills = styles.find(Q('fills'))
        self.fill_id = len(fills)
        fill = ET.SubElement(fills, Q('fill'))
        pattern = ET.SubElement(fill, Q('patternFill'), patternType='solid')
        ET.SubElement(pattern, Q('fgColor'), rgb='FFEAF2F8')
        ET.SubElement(pattern, Q('bgColor'), indexed='64')
        fills.set('count', str(len(fills)))
        self.cache = {}
        self.font_cache = {}

    def derive(self, style):
        if style in self.cache:
            return self.cache[style]
        xf = deepcopy(self.xfs[style])
        font_id = int(xf.get('fontId', '0'))
        if font_id not in self.font_cache:
            font = deepcopy(self.fonts[font_id])
            for tag, attr, val in [('name','val','Arial'),('color','rgb','FF1F4E78'),('i','val','1')]:
                node = font.find(Q(tag))
                if node is None:
                    node = ET.SubElement(font, Q(tag))
                node.attrib.clear()
                node.set(attr,val)
            self.font_cache[font_id] = len(self.fonts)
            self.fonts.append(font)
            self.fonts.set('count', str(len(self.fonts)))
        xf.set('fontId', str(self.font_cache[font_id]))
        xf.set('fillId', str(self.fill_id))
        for attr in ('applyFont','applyFill','applyAlignment'):
            xf.set(attr,'1')
        alignment = xf.find(Q('alignment'))
        if alignment is None:
            alignment = ET.Element(Q('alignment'))
            xf.insert(0, alignment)
        alignment.set('vertical','center')
        alignment.set('wrapText','1')
        alignment.set('shrinkToFit','0')
        self.cache[style] = len(self.xfs)
        self.xfs.append(xf)
        self.xfs.set('count',str(len(self.xfs)))
        return self.cache[style]


def build(package, manifest):
    check_source(package)
    translations = decisions(package, manifest)
    retained = {u['id'] for u in manifest['translation_units'] if u['status'] == 'retain'}
    for item in manifest['occurrences']:
        if item['translation_unit_id'] in retained:
            translations.pop((item['sheet'],item['address']), None)
    styles = TranslationStyles(package.styles)
    patches, expanded, changed_sheets = {}, 0, []
    masters = {}
    for (sheet, cell_address), cell in package.cells.items():
        formula = cell.find(Q('f'))
        if formula is not None and formula.get('t') == 'shared' and formula.text:
            masters[sheet,formula.get('si')] = (cell_address, formula.text)
    for sheet in package.sheets:
        root, name = sheet['root'], sheet['name']
        data = root.find(Q('sheetData'))
        rows = {int(r.get('r')):r for r in data}
        if not any(len(r) for r in rows.values()):
            continue
        max_col = max(coordinates(c.get('r'))[0] for r in rows.values() for c in r)
        max_row = max(rows)
        defaults = root.find(Q('sheetFormatPr'))
        height_default = float(defaults.get('defaultRowHeight','15')) if defaults is not None else 15
        new_rows = []
        for number in range(1,max_row+1):
            original = rows.get(number,ET.Element(Q('row'),r=str(number)))
            source_row = deepcopy(original)
            source_row.set('r',str(number*2-1))
            source_row.attrib.pop('spans',None)
            for cell in source_row:
                old_address = cell.get('r')
                cell.set('r',address(old_address))
                formula = cell.find(Q('f'))
                if formula is not None:
                    text = formula.text
                    if formula.get('t') == 'shared' and not text:
                        origin, master = masters[name,formula.get('si')]
                        text = Translator('='+master,origin=origin).translate_formula(old_address)[1:]
                    formula.attrib.clear()
                    formula.text = map_formula(text)
            target_row = ET.Element(Q('row'), **dict(original.attrib))
            target_row.set('r',str(number*2))
            target_row.attrib.pop('spans',None)
            target_row.attrib.pop('s',None)
            target_row.attrib.pop('customFormat',None)
            height = max(24, float(original.get('ht',str(height_default))))
            originals = {coordinates(c.get('r'))[0]:c for c in original}
            for column in range(1,max_col+1):
                old_address = f'{get_column_letter(column)}{number}'
                source_cell = originals.get(column)
                style_id = int(source_cell.get('s',original.get('s','0'))) if source_cell is not None else int(original.get('s','0'))
                target_cell = ET.SubElement(target_row,Q('c'),r=address(old_address,True),s=str(styles.derive(style_id)))
                text = translations.get((name,old_address))
                if text is not None:
                    target_cell.set('t','inlineStr')
                    inline = ET.SubElement(target_cell,Q('is'))
                    ET.SubElement(inline,Q('t'),{XML_SPACE:'preserve'}).text = text
                    width, _ = layout_geometry(root,old_address)
                    font = styles.fonts[int(styles.xfs[style_id].get('fontId','0'))]
                    sz = font.find(Q('sz'))
                    size = float(sz.get('val','11')) if sz is not None else 11
                    height = max(height,translated_height(text,width,size,font.find(Q('b')) is not None))
            if height > 409.5:
                raise ValueError(f'translation exceeds maximum row height: {name}!{number}; shorten wording or widen the source column')
            target_row.set('ht',str(height))
            target_row.set('customHeight','1')
            expanded += height > 24
            new_rows.extend([source_row,target_row])
        data[:] = new_rows
        dimension = root.find(Q('dimension'))
        if dimension is not None:
            dimension.set('ref',f'A1:{get_column_letter(max_col)}{max_row*2}')
        merges = root.find(Q('mergeCells'))
        if merges is not None:
            originals = list(merges)
            merges[:] = []
            for node in originals:
                first,last = node.get('ref').split(':')
                for translated in (False,True):
                    ET.SubElement(merges,Q('mergeCell'),ref=f'{address(first,translated)}:{address(last,translated)}')
            merges.set('count',str(len(merges)))
        for node in root.iter():
            if node.tag in {Q('conditionalFormatting'),Q('dataValidation'),Q('selection')} and node.get('sqref'):
                node.set('sqref',' '.join(map_reference(x) for x in node.get('sqref').split()))
            if node.tag in {Q('formula'),Q('formula1'),Q('formula2')} and node.text:
                node.text = map_formula(node.text)
            for attr in ('activeCell','topLeftCell'):
                if node.get(attr): node.set(attr,address(node.get(attr)))
            if node.tag == Q('pane') and node.get('state') in {'frozen','frozenSplit'} and node.get('ySplit'):
                node.set('ySplit',str(int(float(node.get('ySplit')))*2))
            if node.tag == Q('hyperlink') and node.get('ref'):
                node.set('ref',map_reference(node.get('ref')))
        for brk in root.findall(f"{Q('rowBreaks')}/{Q('brk')}"):
            brk.set('id',str(int(brk.get('id'))*2))
        setup = root.find(Q('pageSetup'))
        if setup is None or not any(setup.get(k) for k in ('scale','fitToWidth')):
            # Keep orientation and paper size; supply a width fit only when unspecified.
            if setup is None:
                setup = ET.Element(Q('pageSetup'))
                order_after = {Q(x) for x in ('headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst')}
                root.insert(next((i for i,n in enumerate(root) if n.tag in order_after),len(root)),setup)
            setup.set('fitToWidth','1')
            setup.set('fitToHeight','0')
            props = root.find(Q('sheetPr'))
            if props is None:
                props = ET.Element(Q('sheetPr'))
                root.insert(0,props)
            page_props = props.find(Q('pageSetUpPr'))
            if page_props is None: page_props = ET.SubElement(props,Q('pageSetUpPr'))
            page_props.set('fitToPage','1')
        patches[sheet['part']] = serialize(root)
        changed_sheets.append(name)
    workbook = xml(package.data['xl/workbook.xml'])
    for node in workbook.findall(f"{Q('definedNames')}/{Q('definedName')}"):
        if node.text:
            if node.get('name') in {'_xlnm.Print_Area','_xlnm.Print_Titles'}:
                tokens = Tokenizer('='+node.text).items
                node.text = ''.join(map_reference(t.value,True) if t.type == 'OPERAND' and t.subtype == 'RANGE' else t.value for t in tokens)
            else:
                node.text = map_formula(node.text)
    patches['xl/workbook.xml'] = serialize(workbook)
    patches['xl/styles.xml'] = serialize(package.styles)
    if 'xl/calcChain.xml' in package.data:
        chain = xml(package.data['xl/calcChain.xml'])
        for node in chain:
            if node.get('r'): node.set('r',address(node.get('r')))
        patches['xl/calcChain.xml'] = serialize(chain)
    return patches, dict(changed_sheets=changed_sheets,expandedRows=expanded,compressedRows=0)


def verify(source, output, manifest):
    before, after = Package(source), Package(output)
    patches, _ = build(before, manifest)
    errors = []
    if before.data.keys() != after.data.keys():
        errors.append('package-parts-change')
    for part,data in before.data.items():
        actual = after.data.get(part)
        expected = patches.get(part,data)
        if actual != expected:
            if actual is None or part not in patches or canonical(xml(actual)) != canonical(xml(expected)):
                errors.append(f'bilingual-preservation-change:{part}')
    return dict(passed=not errors,errors=errors,checked_cells=len(before.cells),
                preserved_parts=len(before.data)-len(patches),warnings=[])


def apply(source, output, manifest):
    source,output = Path(source).resolve(),Path(output).resolve()
    if source == output:
        raise ValueError('output must not overwrite source')
    package = Package(source)
    patches,report = build(package,manifest)
    output.parent.mkdir(parents=True,exist_ok=True)
    fd,name = tempfile.mkstemp(suffix='.xlsx',dir=output.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        with ZipFile(source) as original, ZipFile(temporary,'w') as target:
            target.comment = original.comment
            for entry in original.infolist():
                target.writestr(entry,patches.get(entry.filename,package.data[entry.filename]))
        check = verify(source,temporary,manifest)
        if not check['passed']:
            raise ValueError(str(check['errors']))
        os.replace(temporary,output)
    finally:
        temporary.unlink(missing_ok=True)
    return report
