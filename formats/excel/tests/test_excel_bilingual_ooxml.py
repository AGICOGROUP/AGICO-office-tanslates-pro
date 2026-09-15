from copy import copy
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from excel_bilingual_ooxml import apply, verify, map_formula, translated_height
from excel_native_ooxml import Package, Q, xml, serialize


class BilingualPreservationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / 'source.xlsx'
        self.output = Path(self.directory.name) / 'output.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = 'Loads'
        ws['A1'] = '设备'
        ws['A2'] = '001'
        ws['B2'] = 1.25
        ws['C2'] = '=B2*2'
        ws['B3'] = 3
        ws['C3'] = '=B3*2'
        ws['A1'].font = Font(name='SimSun',size=16,bold=True,color='FF112233')
        ws['A1'].fill = PatternFill('solid',fgColor='FFFF9900')
        ws['A1'].border = Border(bottom=Side(style='double',color='FF000000'))
        ws['A1'].alignment = Alignment(horizontal='center',vertical='center')
        ws['B2'].number_format = '0.00;[Red]0.00'
        ws.merge_cells('A1:C1')
        ws.row_dimensions[1].height = 28
        ws.column_dimensions['A'].width = 9
        ws.freeze_panes = 'A2'
        ws.print_area = 'A1:C3'
        ws.print_title_rows = '1:1'
        ws.page_setup.orientation = 'landscape'
        ws.oddHeader.center.text = 'Project'
        wb.create_sheet('Empty')
        wb.save(self.source)
        self.manifest = {'output_mode':'bilingual','occurrences':[
            {'id':'Loads!A1','sheet':'Loads','address':'A1','source':'设备','translation_unit_id':'one'},
            {'id':'Loads!A2','sheet':'Loads','address':'A2','source':'001','translation_unit_id':'two'}],
            'translation_units':[
                {'id':'one','source':'设备','translation':'Equipment descriptions and operating requirements','status':'translated'},
                {'id':'two','source':'001','translation':'001','status':'retain'}]}

    def test_source_styles_and_print_metadata_survive_pairing(self):
        apply(self.source,self.output,self.manifest)
        old,new = load_workbook(self.source),load_workbook(self.output)
        a,b = old.active,new.active
        for source,target in [('A1','A1'),('B2','B3'),('C3','C5')]:
            for prop in ['font','fill','border','alignment','number_format','protection']:
                self.assertEqual(copy(getattr(a[source],prop)),copy(getattr(b[target],prop)),prop)
        self.assertEqual(b['C3'].value,'=B3*2')
        self.assertEqual(b['C5'].value,'=B5*2')
        self.assertIsNone(b['A4'].value)
        self.assertIsNone(b['B4'].value)
        self.assertEqual(b['A2'].fill.fgColor.rgb,'FFEAF2F8')
        self.assertEqual(b['A2'].border.bottom.style,'double')
        self.assertGreater(b.row_dimensions[2].height,28)
        self.assertEqual(b['A1'].font.sz,16)
        self.assertEqual(b.column_dimensions['A'].width,9)
        self.assertEqual(b.freeze_panes,'A3')
        self.assertEqual(b.print_title_rows,'$1:$2')
        self.assertIn('$A$1:$C$6',b.print_area)
        self.assertEqual(b.page_setup.orientation,'landscape')
        self.assertEqual(b.page_setup.fitToWidth,1)
        self.assertEqual(b.oddHeader.center.text,'Project')
        self.assertEqual(set(map(str,b.merged_cells)),{'A1:C1','A2:C2'})
        self.assertEqual(new['Empty'].max_row,1)

    def test_shared_formula_followers_remain_formulas(self):
        with ZipFile(self.source) as archive:
            entries = {name:archive.read(name) for name in archive.namelist()}
        root = xml(entries['xl/worksheets/sheet1.xml'])
        cells = {c.get('r'):c for c in root.iter(Q('c'))}
        cells['C2'].find(Q('f')).attrib.update({'t':'shared','si':'0','ref':'C2:C3'})
        follower = cells['C3'].find(Q('f'))
        follower.attrib.update({'t':'shared','si':'0'})
        follower.text = None
        entries['xl/worksheets/sheet1.xml'] = serialize(root)
        with ZipFile(self.source,'w') as archive:
            for name,data in entries.items(): archive.writestr(name,data)
        apply(self.source,self.output,self.manifest)
        ws = load_workbook(self.output).active
        self.assertEqual(ws['C3'].value,'=B3*2')
        self.assertEqual(ws['C5'].value,'=B5*2')
        self.assertIsNone(ws['C6'].value)

    def test_verification_rejects_lost_style_and_numeric_formula_replacement(self):
        apply(self.source,self.output,self.manifest)
        with ZipFile(self.output) as archive:
            entries = {name:archive.read(name) for name in archive.namelist()}
        root = xml(entries['xl/worksheets/sheet1.xml'])
        cells = {c.get('r'):c for c in root.iter(Q('c'))}
        cells['A1'].set('s','0')
        cells['C5'].remove(cells['C5'].find(Q('f')))
        entries['xl/worksheets/sheet1.xml'] = serialize(root)
        with ZipFile(self.output,'w') as archive:
            for name,data in entries.items(): archive.writestr(name,data)
        self.assertFalse(verify(self.source,self.output,self.manifest)['passed'])

    def test_formula_mapping_preserves_literals_functions_and_absolute_refs(self):
        self.assertEqual(map_formula('SUM($A$1:A3)+LOG10(B2)+\'A1 data\'!B4'),
                         'SUM($A$1:A5)+LOG10(B3)+\'A1 data\'!B7')
        self.assertEqual(map_formula('IF(A2="A1",1,0)'), 'IF(A3="A1",1,0)')
        with self.assertRaises(ValueError): map_formula('INDIRECT("A2")')

    def test_narrow_long_labels_get_space_for_every_wrapped_line(self):
        self.assertGreater(translated_height('Tangent of Power Factor Angle (tanφ)',5.25,12),150)
        self.assertGreater(translated_height('Belt Conveyor No. 1',13.25,12),40)


if __name__ == '__main__': unittest.main()
