"""Production native writer contracts: fidelity, coverage and failure recovery."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "excel_native_ooxml.py"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def fixture(path, rows=0):
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        "xl/workbook.xml": f'<workbook xmlns="{NS}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Plant" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/sharedStrings.xml": f'<sst xmlns="{NS}" count="2" uniqueCount="1"><si><r><rPr><b/></rPr><t>设备</t></r><r><rPr><i/></rPr><t>名称</t></r></si></sst>',
        "xl/styles.xml": f'<styleSheet xmlns="{NS}"><fonts count="1"><font><sz val="11"/></font></fonts><fills count="1"><fill/></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf/></cellStyleXfs><cellXfs count="1"><xf fontId="0" fillId="0" borderId="0"/></cellXfs></styleSheet>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{NS}"><sheetFormatPr defaultRowHeight="15"/><cols><col min="1" max="3" width="10" customWidth="1"/></cols><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>0</v></c><c r="C1"><v>45</v></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>运行</t></is></c><c r="B2"><f>COUNTIF(A2,"运行")</f><v>1</v></c></row><row r="3" ht="30" customHeight="1"><c r="A3" t="inlineStr"><is><t>备注</t></is></c></row>' + ''.join(f'<row r="{r}"><c r="A{r}" t="inlineStr"><is><t>设备 {r}</t></is></c><c r="B{r}"><v>{r}</v></c></row>' for r in range(4, rows + 4)) + '</sheetData><mergeCells count="1"><mergeCell ref="A3:B3"/></mergeCells><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>',
        "xl/charts/chart1.xml": '<chart xmlns="http://schemas.openxmlformats.org/drawingml/2006/chart"><title>原始标题</title></chart>',
        "xl/comments/comment1.xml": f'<comments xmlns="{NS}"><commentList><comment ref="A1"><text><t>原始批注</t></text></comment></commentList></comments>',
        "xl/worksheets/_rels/sheet1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="external" TargetMode="External" Target="https://example.invalid/data" Type="externalLink"/></Relationships>',
        "xl/drawings/drawing1.xml": '<drawing>unknown extension text</drawing>',
        "customXml/item1.xml": '<unknown preserve="exactly" />',
    }
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content.encode() if isinstance(content, str) else content)
    return parts


def manifest(path, source):
    cells = [("A1", "设备名称", "Equipment name with extended description"),
             ("B1", "设备名称", "Name"), ("A2", "运行", "运行"), ("A3", "备注", "Notes")]
    result = {"source_file": str(source), "occurrences": [], "translation_units": []}
    for address, text, target in cells:
        result["occurrences"].append({"id": f"Plant!{address}", "sheet": "Plant", "address": address,
                                      "kind": "cell", "source": text, "translation_unit_id": address,
                                      "formula_dependency": address == "A2"})
        result["translation_units"].append({"id": address, "source": text, "translation": target,
                                           "status": "retain" if text == target else "translated"})
    path.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return result


class NativeWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source, self.output, self.manifest = [self.root / name for name in ("source.xlsx", "out.xlsx", "manifest.json")]
        fixture(self.source)
        self.data = manifest(self.manifest, self.source)

    def run_native(self, action, ok=True):
        command = [sys.executable, str(SCRIPT), action, "--source", str(self.source)]
        if action != "inspect":
            command += ["--output", str(self.output), "--manifest", str(self.manifest)]
        result = subprocess.run(command, capture_output=True, encoding="utf-8")
        if ok:
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            return json.loads(result.stdout)
        self.assertNotEqual(result.returncode, 0)
        return result

    def test_rich_shared_strings_split_and_all_other_parts_survive(self):
        report = self.run_native("apply")
        self.assertEqual(report["changed_sheets"], ["Plant"])
        with ZipFile(self.source) as source, ZipFile(self.output) as output:
            self.assertEqual(source.namelist(), output.namelist())
            for name in source.namelist():
                if name not in {"xl/worksheets/sheet1.xml", "xl/styles.xml"}:
                    self.assertEqual(source.read(name), output.read(name), name)
            sheet = ET.fromstring(output.read("xl/worksheets/sheet1.xml"))
            for address, expected in [("A1", "Equipment name with extended description"), ("B1", "Name")]:
                cell = sheet.find(f'.//{{{NS}}}c[@r="{address}"]')
                self.assertEqual(cell.get("t"), "inlineStr")
                self.assertEqual("".join(cell.itertext()), expected)
                self.assertIsNotNone(cell.find(f'.//{{{NS}}}b'))
                self.assertIsNotNone(cell.find(f'.//{{{NS}}}i'))
            self.assertGreater(float(sheet.find(f'.//{{{NS}}}row[@r="1"]').get("ht")), 15)
            self.assertIsNone(sheet.find(f'.//{{{NS}}}row[@r="2"]').get("ht"))
            self.assertEqual(sheet.find(f'.//{{{NS}}}row[@r="3"]').get("ht"), "30")
        self.assertTrue(self.run_native("verify")["passed"])

    def test_inspection_discloses_complex_parts_and_operational_text(self):
        report = self.run_native("inspect")
        self.assertEqual(len(report["sheets"]), 1)
        cell = next(c for c in report["occurrences"] if c["address"] == "A2")
        self.assertTrue(cell["formula_dependency"])
        self.assertTrue(any("chart1.xml" in w for w in report["warnings"]))
        self.assertTrue(any("comment1.xml" in w for w in report["warnings"]))

    def formula_fixture(self, formula, shared=False):
        parts = fixture(self.source)
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"].encode())
        node = root.find(f'.//{{{NS}}}f')
        node.text = formula
        cell = root.find(f'.//{{{NS}}}c[@r="C1"]')
        cell.set("t", "inlineStr")
        cell.clear(keep_tail=True)
        cell.set("r", "C1")
        cell.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(cell, f"{{{NS}}}is"), f"{{{NS}}}t").text = "输入"
        if shared:
            node.set("t", "shared")
            node.set("si", "0")
            node.set("ref", "B2:B3")
            follower = ET.SubElement(root.find(f'.//{{{NS}}}row[@r="3"]'), f"{{{NS}}}c", r="B3")
            ET.SubElement(follower, f"{{{NS}}}f", t="shared", si="0")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        with ZipFile(self.source, "w", ZIP_DEFLATED) as archive:
            for name, content in parts.items():
                archive.writestr(name, content)

    def test_literal_indirect_retains_only_referenced_input(self):
        self.formula_fixture('SUM(INDIRECT("C1"))')
        report = self.run_native("inspect")
        self.assertEqual({c["address"] for c in report["occurrences"] if c["formula_dependency"]}, {"C1"})

    def test_indirect_explicit_a1_mode_and_quoted_sheet(self):
        self.formula_fixture('SUM(INDIRECT("\'Plant\'!C1",1))')
        report = self.run_native("inspect")
        self.assertEqual({c["address"] for c in report["occurrences"] if c["formula_dependency"]}, {"C1"})

    def test_print_area_does_not_mark_display_labels_as_formula_inputs(self):
        parts = fixture(self.source)
        parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
            '</workbook>', '<definedNames><definedName name="_xlnm.Print_Area">Plant!$A$1:$C$3</definedName>'
            '<definedName name="_xlnm.Print_Titles">Plant!$1:$3</definedName></definedNames></workbook>')
        with ZipFile(self.source, "w", ZIP_DEFLATED) as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        report = self.run_native("inspect")
        self.assertEqual({c["address"] for c in report["occurrences"] if c["formula_dependency"]}, {"A2"})

    def test_static_offset_and_whole_ranges_are_bounded(self):
        for formula, expected in [('SUM(OFFSET(B1,0,1))', {"B1", "C1"}),
                                  ('SUM(C:C)', {"C1"}), ('SUM(3:3)', {"A3"})]:
            with self.subTest(formula=formula):
                self.formula_fixture(formula)
                report = self.run_native("inspect")
                self.assertEqual({c["address"] for c in report["occurrences"] if c["formula_dependency"]}, expected)

    def test_shared_formula_relative_inputs_are_expanded(self):
        self.formula_fixture('A2', shared=True)
        report = self.run_native("inspect")
        self.assertEqual({c["address"] for c in report["occurrences"] if c["formula_dependency"]}, {"A2", "A3"})

    def test_unresolved_references_fail_explicitly_without_replacing_output(self):
        for formula in ['SUM(INDIRECT(A1))', 'SUM(OFFSET(A1,B1,0))', 'SUM(Table1[Label])',
                        'SUM(_xlfn.INDIRECT("C1"))', 'SUM(_xlfn.OFFSET(A1,0,2))']:
            with self.subTest(formula=formula):
                self.formula_fixture(formula)
                self.output.write_bytes(b"previous deliverable")
                result = self.run_native("apply", ok=False)
                self.assertIn("unresolved formula reference", result.stdout)
                self.assertIn("Plant", result.stdout)
                self.assertEqual(self.output.read_bytes(), b"previous deliverable")

    def test_rich_text_inherited_namespaces_pass_but_real_font_change_fails(self):
        # Shared strings and worksheets can have different namespace declarations.
        with ZipFile(self.source) as archive:
            entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
        with ZipFile(self.source, "w") as archive:
            for entry, data in entries:
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<worksheet ", b'<worksheet xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData" ')
                archive.writestr(entry, data)
        self.run_native("apply")
        self.assertTrue(self.run_native("verify")["passed"])
        with ZipFile(self.output) as archive:
            entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
        with ZipFile(self.output, "w") as archive:
            for entry, data in entries:
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<b/>", b'<sz val="30"/>', 1)
                archive.writestr(entry, data)
        self.assertIn("rich-format-change", self.run_native("verify", ok=False).stdout)

    def test_xml_comments_survive_inspect_apply_and_verify(self):
        comment = b"<!-- Generated by reporting tool -->"
        with ZipFile(self.source) as archive:
            entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
        with ZipFile(self.source, "w") as archive:
            for entry, data in entries:
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<sheetData>", comment + b"<sheetData>")
                archive.writestr(entry, data)
        inspection = self.run_native("inspect")
        self.assertEqual(len(inspection["occurrences"]), 4)
        self.run_native("apply")
        self.assertTrue(self.run_native("verify")["passed"])
        with ZipFile(self.output) as archive:
            self.assertEqual(archive.read("xl/worksheets/sheet1.xml").count(comment), 1)

    def test_source_mismatch_leaves_previous_output_untouched(self):
        self.output.write_bytes(b"previous verified deliverable")
        self.data["occurrences"][0]["source"] = "changed source"
        self.manifest.write_text(json.dumps(self.data), encoding="utf-8")
        self.run_native("apply", ok=False)
        self.assertEqual(self.output.read_bytes(), b"previous verified deliverable")

    def test_operational_text_cannot_be_changed_even_with_forged_flag(self):
        self.data["translation_units"][2]["translation"] = "Running"
        self.data["occurrences"][2]["formula_dependency"] = False
        self.manifest.write_text(json.dumps(self.data), encoding="utf-8")
        self.run_native("apply", ok=False)
        self.assertFalse(self.output.exists())

    def test_macro_payload_is_rejected_even_in_renamed_xlsx(self):
        with ZipFile(self.source, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"DO NOT EXECUTE")
        self.assertIn("macro", self.run_native("inspect", ok=False).stdout)

    def test_missing_coverage_does_not_silently_drop_translation(self):
        self.data["occurrences"].pop()
        self.manifest.write_text(json.dumps(self.data), encoding="utf-8")
        self.assertIn("coverage", self.run_native("apply", ok=False).stdout)

    def test_replace_failure_preserves_previous_output_and_cleans_temporary(self):
        sys.path.insert(0, str(SCRIPT.parent))
        spec = importlib.util.spec_from_file_location("excel_native_test", SCRIPT)
        native = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(native)
        self.output.write_bytes(b"previous verified deliverable")
        with patch.object(native.os, "replace", side_effect=OSError("simulated locked output")):
            with self.assertRaisesRegex(OSError, "locked output"):
                native.apply(self.source, self.output, self.data)
        self.assertEqual(self.output.read_bytes(), b"previous verified deliverable")
        self.assertFalse(list(self.root.glob(".excel-native-*")))

    def test_verifier_detects_rich_properties_and_style_association_damage(self):
        self.run_native("apply")
        original = self.output.read_bytes()
        for old, new in [(b"<b/>", b"<strike/>"), (b's="1"', b's="0"')]:
            self.output.write_bytes(original)
            with ZipFile(self.output) as z:
                entries = [(info, z.read(info.filename)) for info in z.infolist()]
            with ZipFile(self.output, "w") as z:
                for info, data in entries:
                    z.writestr(info, data.replace(old, new, 1) if info.filename == "xl/worksheets/sheet1.xml" else data)
            self.run_native("verify", ok=False)

    def test_real_workbook_reopens_with_chart_comment_image_and_formula_intact(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        from openpyxl.drawing.image import Image as ExcelImage
        from PIL import Image
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plant"
        sheet.append(["设备名称", 10])
        sheet.append(["风机", 20, "=SUM(B1:B2)"])
        sheet["A1"].comment = Comment("保留批注", "Engineer")
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
        sheet.add_chart(chart, "E1")
        picture = self.root / "image.png"
        Image.new("RGB", (24, 24), "blue").save(picture)
        sheet.add_image(ExcelImage(str(picture)), "E15")
        workbook.save(self.source)
        units = [("A1", "设备名称", "Equipment name"), ("A2", "风机", "Fan")]
        self.data = {"occurrences": [{"id": address, "kind": "cell", "sheet": "Plant", "address": address,
                                      "source": source, "translation_unit_id": address} for address, source, _ in units],
                     "translation_units": [{"id": address, "source": source, "translation": target,
                                            "status": "translated"} for address, source, target in units]}
        self.manifest.write_text(json.dumps(self.data), encoding="utf-8")
        self.run_native("apply")
        self.assertTrue(self.run_native("verify")["passed"])
        with ZipFile(self.source) as a, ZipFile(self.output) as b:
            for part in a.namelist():
                if part not in {"xl/worksheets/sheet1.xml", "xl/styles.xml"}:
                    self.assertEqual(a.read(part), b.read(part), part)
        reopened = load_workbook(self.output)
        self.assertEqual(reopened.active["A1"].value, "Equipment name")
        self.assertEqual(reopened.active["C2"].value, "=SUM(B1:B2)")
        self.assertEqual(reopened.active["A1"].comment.text, "保留批注")
        self.assertEqual(len(reopened.active._charts), 1)
        self.assertEqual(len(reopened.active._images), 1)
        reopened.close()

    def test_verifier_detects_number_formula_unknown_part_and_structure_damage(self):
        self.run_native("apply")
        original = self.output.read_bytes()
        changes = [("xl/worksheets/sheet1.xml", b">45<", b">46<", "non-text-change"),
                   ("xl/worksheets/sheet1.xml", b"COUNTIF", b"SUMIF", "formula-change"),
                   ("customXml/item1.xml", b"exactly", b"changed", "untouched-part-change"),
                   ("xl/worksheets/sheet1.xml", b'A3:B3', b'A3:C3', "structure-change")]
        for part, old, new, error in changes:
            self.output.write_bytes(original)
            with ZipFile(self.output) as z:
                entries = [(info, z.read(info.filename)) for info in z.infolist()]
            with ZipFile(self.output, "w") as z:
                for info, data in entries:
                    z.writestr(info, data.replace(old, new) if info.filename == part else data)
            result = self.run_native("verify", ok=False)
            self.assertIn(error, result.stdout)


if __name__ == "__main__":
    unittest.main()
