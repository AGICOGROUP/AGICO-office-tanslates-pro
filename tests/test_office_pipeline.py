import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import office_pipeline as office


class OfficePipelineTests(unittest.TestCase):
    def test_incomplete_prepare_cannot_merge_finalize_or_claim_delivery(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, job = self.create_source(root), root / "job"
            office.prepare_job(source, job, "en")
            state_path = job / "office-job.json"
            state = self.read(state_path)
            state["prepared"] = False
            state_path.write_text(json.dumps(state), encoding="utf-8")
            for operation in (lambda: office.merge_job(job),
                              lambda: office.finalize_job(job, root / "out.docx"),
                              lambda: office.status_job(job)):
                with self.subTest(operation=operation):
                    with self.assertRaisesRegex(ValueError, "prepare"):
                        operation()
            self.assertFalse((root / "out.docx").exists())
            # Recovery uses the supported command and needs no hand-edited state.
            office.prepare_job(source, job, "en")
            self.assertEqual(office.status_job(job)["next_stage"], "translate")

    def test_adapter_without_structured_result_is_not_success(self):
        result = office.subprocess.CompletedProcess([], 0, stdout="interrupted before report\n", stderr="")
        with mock.patch.object(office.subprocess, "run", return_value=result):
            with self.assertRaisesRegex(RuntimeError, "result"):
                office.run_adapter("word", "validate", "missing.docx")

    def test_status_does_not_hide_pending_work_behind_stale_delivery(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'source.docx'
            doc = Document(); doc.add_paragraph('BASIC REQUIREMENTS'); doc.save(source)
            job = root / 'job'
            office.prepare_job(source, job, 'zh-CN')
            manifest_path = job / 'translation-manifest.json'
            manifest = self.read(manifest_path)
            output = root / 'out.docx'
            output.write_bytes(source.read_bytes())
            state_path = job / 'office-job.json'
            state = self.read(state_path)
            delivered = {
                'next_stage': 'deliver', 'output': str(output),
                'output_sha256': office.file_hash(output), 'warnings': []}
            state.setdefault('stages', {})['deliver'] = {
                'key': office.fingerprint({'manifest': manifest, 'output': str(output)}),
                'output_sha256': delivered['output_sha256'], 'result': delivered}
            state_path.write_text(json.dumps(state), encoding='utf-8')
            result = office.status_job(job)
            self.assertEqual(result['next_stage'], 'translate')
            self.assertEqual(result['pending_count'], 1)

    def create_source(self, root):
        source = root / "source.docx"
        doc = Document()
        doc.add_paragraph("设备")
        doc.add_paragraph("功率45kW")
        doc.save(source)
        return source

    def read(self, path):
        return json.loads(path.read_text(encoding="utf-8"))

    def test_partial_repair_and_repeat_finalize_reuse_completed_work(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self.create_source(root)
            job = root / "job"
            result = office.prepare_job(source, job, "en")
            self.assertEqual(result["pending_count"], 2)
            workpath = job / "translation-worklist.json"
            work = self.read(workpath)
            work["translation_units"][0]["translation"] = "Equipment"
            work["translation_units"][1]["translation"] = "Power 46 kW"
            workpath.write_text(json.dumps(work), encoding="utf-8")
            report = office.merge_job(job, workpath)
            self.assertEqual(report["pending_count"], 1)
            self.assertEqual(len(self.read(workpath)["translation_units"]), 1)
            self.assertEqual(office.finalize_job(job, root / "out.docx")["next_stage"], "translate")
            work = self.read(workpath)
            work["translation_units"][0]["translation"] = "Power 45 kW"
            workpath.write_text(json.dumps(work), encoding="utf-8")
            with mock.patch.object(office, "run_adapter", wraps=office.run_adapter) as runner:
                first = office.finalize_job(job, root / "out.docx")
                calls = runner.call_count
                again = office.finalize_job(job, root / "out.docx")
                self.assertEqual(runner.call_count, calls)
            self.assertEqual(first["next_stage"], "deliver")
            self.assertEqual(first["output_sha256"], again["output_sha256"])
            self.assertEqual([p.text for p in Document(root / "out.docx").paragraphs], ["Equipment", "Power 45 kW"])

    def test_repeated_prepare_preserves_unsubmitted_worklist(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self.create_source(root)
            job = root / "job"
            office.prepare_job(source, job, "en")
            workpath = job / "translation-worklist.json"
            work = self.read(workpath)
            work["translation_units"][0]["translation"] = "Equipment"
            workpath.write_text(json.dumps(work), encoding="utf-8")
            with mock.patch.object(office, "run_adapter", side_effect=AssertionError("must resume")):
                office.prepare_job(source, job, "en")
            self.assertEqual(self.read(workpath)["translation_units"][0]["translation"], "Equipment")
            with self.assertRaisesRegex(ValueError, "identity"):
                office.prepare_job(source, job, "zh")

    def test_changed_source_and_original_output_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self.create_source(root)
            job = root / "job"
            office.prepare_job(source, job, "en")
            with self.assertRaisesRegex(ValueError, "overwrite"):
                office.finalize_job(job, source)
            source.write_bytes(b"changed source")
            with self.assertRaisesRegex(ValueError, "source.*changed"):
                office.finalize_job(job, root / "out.docx")

    def test_ppt_available_slides_return_review_and_resume_without_rerender(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "设备"
            deck.save(source)
            job = root / "job"
            office.prepare_job(source, job, "en")
            workpath = job / "translation-worklist.json"
            work = self.read(workpath)
            work["translation_units"][0]["translation"] = "Equipment"
            workpath.write_text(json.dumps(work), encoding="utf-8")
            actual_run = office.run_adapter
            calls = []
            def run(format_name, *arguments):
                calls.append(arguments[0])
                if arguments[0] == "render":
                    return {"target": {"slides_rendered": 1, "render_directory": str(job / "renders")}}
                if arguments[0] == "deliver":
                    self.assertIn("--visual-review-passed", arguments)
                    return {"warnings": []}
                return actual_run(format_name, *arguments)
            with mock.patch.object(office, "run_adapter", side_effect=run):
                first = office.finalize_job(job, root / "out.pptx")
                self.assertEqual(first["next_stage"], "visual-review")
                final = office.finalize_job(job, root / "out.pptx", visual_review_passed=True)
                self.assertEqual(final["next_stage"], "deliver")
            self.assertEqual(calls.count("apply"), 1)
            self.assertEqual(calls.count("render"), 1)

    def test_failed_native_overlay_returns_only_affected_image_for_repair(self):
        from io import BytesIO
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, job = root / "source.pptx", root / "job"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            data = BytesIO()
            Image.new("RGB", (40, 40), "white").save(data, "PNG")
            data.seek(0)
            host = slide.shapes.add_picture(data, Inches(1), Inches(1), Inches(2), Inches(2))
            deck.save(source)
            office.prepare_job(source, job, "en")
            workpath = job / "translation-worklist.json"
            work = self.read(workpath)
            work["images"][0].update(decision="overlay", overlays=[{
                "id": "bad-font", "translation": "Cooler", "source_text": "冷却机",
                "location": {"page_or_slide": 1, "host_shape_id": host.shape_id},
                "source_region": {"x": .1, "y": .1, "w": .5, "h": .1},
                "region": {"x": .1, "y": .2, "w": .5, "h": .1},
                "style": {"font_name": "Arial", "font_size_pt": 0, "bold": False, "text_rgb": "000000", "align": "left"}}])
            workpath.write_text(json.dumps(work), encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "bad-font"):
                office.finalize_job(job, root / "out.pptx")
            retry = self.read(workpath)
            self.assertEqual(len(retry["images"]), 1)
            self.assertIn("font_size_pt", retry["images"][0]["decision_error"])
            self.assertEqual(retry["images"][0]["overlays"][0]["id"], "bad-font")
            retry["images"][0]["overlays"][0]["style"]["font_size_pt"] = 12
            workpath.write_text(json.dumps(retry), encoding="utf-8")
            self.assertTrue(office.merge_job(job)["ready"])


if __name__ == "__main__":
    unittest.main()
