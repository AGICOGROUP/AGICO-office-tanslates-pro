"""Public workflow regressions: bounded reading and verified publication."""
import json
import os
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest
from unittest import mock

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import office_pipeline as office
from translation_core import build_worklist, merge_decisions


class OfficeUpgradeTests(unittest.TestCase):
    def test_cli_batch_unicode_survives_a_legacy_console_encoding(self):
        with tempfile.TemporaryDirectory() as directory:
            _, job = self.prepare(Path(directory))
            result = subprocess.run(
                [sys.executable, str(Path(office.__file__)), 'batch', '--job-dir', str(job)],
                env={**os.environ, 'PYTHONIOENCODING': 'gbk'},
                capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))
            report = json.loads(result.stdout.decode('utf-8'))
            self.assertEqual(report['translation_units'][0]['source'], '冷却机功率45kW')

    def test_explicit_retention_can_omit_duplicate_source_text_but_not_reason(self):
        manifest = {'target_language':'zh-CN', 'units':[{'id':1, 'source':'BEUMER', 'target':''}]}
        identity = build_worklist(manifest, 'word')['job_identity']
        decision = {'id':1, 'status':'retain', 'reason':'Equipment manufacturer brand name'}
        merged, report = merge_decisions(manifest, {'job_identity':identity, 'translation_units':[decision]}, 'word')
        self.assertTrue(report['ready'])
        self.assertEqual(merged['units'][0]['target'], 'BEUMER')
        del decision['reason']
        self.assertFalse(merge_decisions(manifest, {'job_identity':identity, 'translation_units':[decision]}, 'word')[1]['ready'])

    def prepare(self, root, batch_chars=30):
        source = root / 'source.docx'
        doc = Document()
        doc.add_paragraph('冷却机功率45kW')
        doc.add_paragraph('水泥磨处理能力100t/h')
        doc.save(source)
        job = root / 'job'
        office.prepare_job(source, job, 'en', batch_chars=batch_chars)
        return source, job

    def decisions(self, job):
        work = office.read_json(job / 'translation-worklist.json')
        wording = {'冷却机功率45kW': 'Cooler power 45kW',
                   '水泥磨处理能力100t/h': 'Cement mill capacity 100t/h'}
        for item in work['translation_units']:
            item['translation'] = wording[item['source']]
        path = job / 'decisions.json'
        office.atomic_json(path, work)
        return path

    def test_batch_exports_only_selected_pending_units_and_matching_terminology(self):
        with tempfile.TemporaryDirectory() as directory:
            source, job = self.prepare(Path(directory))
            original = (job / 'translation-worklist.json').read_bytes()
            data = office.batch_job(job, index=1)
            self.assertEqual([u['source'] for u in data['translation_units']], ['冷却机功率45kW'])
            self.assertEqual(data['batch']['remaining_batches'], 1)
            terms = [entry['source'] for entry in data['glossary']['entries']]
            self.assertIn('冷却机', terms)
            self.assertNotIn('水泥磨', terms)
            self.assertNotIn('text_indices', data['glossary']['entries'][0])
            self.assertEqual((job / 'translation-worklist.json').read_bytes(), original)
            data['translation_units'][0]['translation'] = 'Cooler power 45kW'
            path = job / 'first.json'
            office.atomic_json(path, data)
            self.assertEqual(office.merge_job(job, path)['pending_count'], 1)
            self.assertEqual(office.batch_job(job)['translation_units'][0]['source'], '水泥磨处理能力100t/h')

    def test_batch_refuses_to_overwrite_original_or_pipeline_owned_state(self):
        with tempfile.TemporaryDirectory() as directory:
            source, job = self.prepare(Path(directory))
            for output in [source, job / 'translation-manifest.json', job / 'office-job.json']:
                before = output.read_bytes()
                with self.assertRaises(ValueError):
                    office.batch_job(job, output=output)
                self.assertEqual(output.read_bytes(), before)

    def test_validation_failure_never_overwrites_previous_deliverable(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, job = self.prepare(root)
            decisions = self.decisions(job)
            output = root / 'translated.docx'
            output.write_bytes(b'previous deliverable')
            real = office.run_adapter
            def fail_validation(fmt, *args):
                if args[0] == 'validate':
                    raise RuntimeError('injected content damage')
                return real(fmt, *args)
            with mock.patch.object(office, 'run_adapter', side_effect=fail_validation):
                with self.assertRaisesRegex(RuntimeError, 'content damage'):
                    office.finalize_job(job, output, decisions=decisions)
            self.assertEqual(output.read_bytes(), b'previous deliverable')
            result = office.finalize_job(job, output)
            self.assertEqual(result['next_stage'], 'deliver')
            self.assertEqual(result['output'], str(output.resolve()))
            self.assertEqual(office.status_job(job)['output'], str(output.resolve()))
            self.assertEqual([p.text for p in Document(output).paragraphs],
                             ['Cooler power 45kW', 'Cement mill capacity 100t/h'])
            self.assertEqual(result['output_sha256'], office.file_hash(output))


if __name__ == '__main__':
    unittest.main()
