import copy
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import office_pipeline as office
from translation_core import build_worklist, merge_decisions
from translation_quality import parameter_mismatch


class TranslationEfficiencyTests(unittest.TestCase):
    def test_fresh_word_job_delivers_compact_translations_without_touching_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, job, output = root / "source.docx", root / "job", root / "zh.docx"
            doc = Document()
            doc.add_paragraph("connected to6.3 KV switchboards", style="Heading 1")
            doc.add_paragraph("medium voltage6.3KV distribution\tMotor ab123")
            doc.save(source)
            source_hash = office.file_hash(source)
            office.prepare_job(source, job, "zh-CN")
            work = office.read_json(job / "translation-worklist.json")
            decisions = {"job_identity": work["job_identity"], "translation_units": [
                {"id": 1, "translation": "连接至6.3 KV开关柜"},
                {"id": 2, "translation": "中压6.3KV配电\t电动机ab123"}]}
            path = job / "decisions.json"
            office.atomic_json(path, decisions)
            result = office.finalize_job(job, output, decisions=path)
            self.assertEqual(result["next_stage"], "deliver")
            translated = Document(output)
            self.assertEqual([p.text for p in translated.paragraphs], [
                "连接至6.3 KV开关柜", "中压6.3KV配电\t电动机ab123"])
            self.assertEqual(translated.paragraphs[0].style.name, "Heading 1")
            self.assertEqual(office.file_hash(source), source_hash)

    def test_glued_voltage_prose_translates_without_preserving_english(self):
        for source, target in [
            ("connected to6.3 KV switchboards", "连接至6.3 KV开关柜"),
            ("medium voltage6.3KV distribution", "中压6.3KV配电"),
            ("VOLTAGE6.3KV", "电压6.3KV"),
        ]:
            with self.subTest(source=source):
                self.assertFalse(parameter_mismatch(source, target))

    def test_glued_prose_still_protects_value_unit_and_real_models(self):
        for source, target in [
            ("connected to6.3 KV", "连接至6.4 KV"),
            ("voltage6.3kV", "电压6.3V"),
            ("voltage6.3kV", "电压6.3KV"),
            ("Motor ab123", "电动机ab124"),
            ("Motor To6.3X", "电动机6.3X"),
            ("voltage6.3KV and 6.3KV", "电压6.3KV"),
            ("Tolerance ±5 mm", "公差5 mm"),
        ]:
            with self.subTest(source=source):
                self.assertTrue(parameter_mismatch(source, target))

    def test_reference_hyphens_are_not_numeric_signs(self):
        self.assertFalse(parameter_mismatch("Annexure-15, §12.17", "附件15，第12.17节"))
        self.assertTrue(parameter_mismatch("Annexure-15, §12.17", "附件15，第12.18节"))
        self.assertFalse(parameter_mismatch("水分-0.57%", "Moisture-0.57%"))
        self.assertTrue(parameter_mismatch("Temperature -10°C", "温度10°C"))

    def test_compact_decisions_use_job_identity_and_reject_explicit_source_conflicts(self):
        for fmt, collection, src, dst in [
            ("word", "units", "source", "target"),
            ("excel", "translation_units", "source", "translation"),
            ("ppt", "translation_units", "source_text", "translation"),
        ]:
            with self.subTest(format=fmt):
                manifest = {"target_language": "zh-CN", collection: [
                    {"id": 1, src: "Power 45kW", dst: ""}]}
                work = {"job_identity": build_worklist(manifest, fmt)["job_identity"],
                        "translation_units": [{"id": 1, "translation": "功率45kW"}]}
                merged, report = merge_decisions(manifest, work, fmt)
                self.assertTrue(report["ready"])
                self.assertEqual(merged[collection][0][dst], "功率45kW")
                work["translation_units"][0]["source"] = "Power 99kW"
                self.assertFalse(merge_decisions(manifest, work, fmt)[1]["ready"])
                work["job_identity"] = "wrong-job"
                with self.assertRaises(ValueError):
                    merge_decisions(manifest, work, fmt)

    def test_word_batches_remove_only_neighbors_visible_in_same_batch(self):
        first, second = "First paragraph " * 20, "Second paragraph " * 20
        manifest = {"units": [
            {"id": 1, "source": first, "target": "", "context": {
                "part": "word/document.xml", "style": "Heading1", "before": "External heading", "after": second[:160]}},
            {"id": 2, "source": second, "target": "", "context": {
                "part": "word/document.xml", "style": "Body", "before": first[-160:], "after": "External end"}},
        ]}
        saved = copy.deepcopy(manifest)
        work = build_worklist(manifest, "word", max_chars=1000)
        self.assertEqual(len(work["batches"]), 1)
        self.assertNotIn("after", work["translation_units"][0]["context"])
        self.assertNotIn("before", work["translation_units"][1]["context"])
        self.assertEqual(work["translation_units"][0]["context"]["before"], "External heading")
        self.assertEqual(work["translation_units"][1]["context"]["after"], "External end")
        self.assertEqual(work["translation_units"][0]["context"]["style"], "Heading1")
        split = build_worklist(manifest, "word", max_chars=650)
        self.assertEqual(len(split["batches"]), 2)
        self.assertEqual(split["translation_units"][0]["context"]["after"], second[:160])
        self.assertEqual(split["translation_units"][1]["context"]["before"], first[-160:])
        self.assertEqual(manifest, saved)

    def test_unrelated_neighbors_and_other_parts_keep_context(self):
        manifest = {"units": [
            {"id": 1, "source": "Motor", "target": "", "context": {"part": "a", "after": "Drive"}},
            {"id": 2, "source": "Drive", "target": "", "context": {"part": "b", "before": "Motor"}},
        ]}
        work = build_worklist(manifest, "word")
        self.assertEqual(work["translation_units"][0]["context"]["after"], "Drive")
        self.assertEqual(work["translation_units"][1]["context"]["before"], "Motor")


if __name__ == "__main__":
    unittest.main()
